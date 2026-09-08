"""Evaluate every formula in the time card workbooks in Python (SUM ranges, additions, the MATCH test),
print the tie-out, and check the text for anything that must not leave the desk."""
import re
import sys
import os

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
FORBIDDEN = ['dev390397', 'ZK5LG9V', 'service-now.com', 'x_196061', 'bofasim', 'Claude', 'Anthropic', 'GPT', 'OpenAI',
             'LLM', ' AI ', 'PDI', 'snui', 'stand-in']


def evaluate(ws):
    """Return {ref: value} with formulas resolved. Supports =SUM(A1:A5), =A1+B1+..., =IF(AND(a=b,...),"x","y")."""
    vals = {}

    def val(ref):
        if ref in vals:
            return vals[ref]
        v = ws[ref].value
        if isinstance(v, str) and v.startswith('='):
            v = calc(v[1:])
        vals[ref] = 0 if v is None else v
        return vals[ref]

    def expr(e):
        e = e.strip()
        m = re.fullmatch(r'SUM\(([A-Z]+)(\d+):([A-Z]+)(\d+)\)', e)
        if m:
            col, r1, r2 = m.group(1), int(m.group(2)), int(m.group(4))
            return sum(val('%s%d' % (col, r)) for r in range(r1, r2 + 1))
        if '+' in e:
            return sum(expr(p) for p in e.split('+'))
        if re.fullmatch(r'[A-Z]+\d+', e):
            return val(e)
        return float(e)

    def calc(f):
        m = re.fullmatch(r'IF\(AND\((.*)\),"(\w+)","(\w+)"\)', f)
        if m:
            conds = m.group(1).split(',')
            ok = all(abs(expr(a) - expr(b)) < 1e-9 for a, b in (c.split('=') for c in conds))
            return m.group(2) if ok else m.group(3)
        return expr(f)

    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith('='):
                vals[c.coordinate] = calc(c.value[1:])
    return vals


def main(paths):
    bad = 0
    for p in paths:
        wb = load_workbook(p)
        print('=====', os.path.basename(p))
        text = ' '.join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None)
        hits = [w for w in FORBIDDEN if w.lower() in text.lower()]
        print('scrub hits:', hits or 'none')
        bad += len(hits)
        tc = evaluate(wb['ServiceNow Daily Timecard'])
        rc = evaluate(wb['Deloitte Reconciliation'])
        ws = wb['Deloitte Reconciliation']
        for r in range(6, ws.max_row + 1):
            if ws['J%d' % r].value:
                label = ws['A%d' % r].value or ''
                print('%-14s %-4s Deloitte %5.2f + %5.2f = %5.2f | SN %5.2f + %5.2f + %5.2f = %5.2f | %s' % (
                    label, ws['B%d' % r].value or '', rc.get('C%d' % r, ws['C%d' % r].value or 0), rc.get('D%d' % r, ws['D%d' % r].value or 0),
                    rc['E%d' % r], rc.get('F%d' % r, ws['F%d' % r].value or 0), rc.get('G%d' % r, ws['G%d' % r].value or 0),
                    rc.get('H%d' % r, ws['H%d' % r].value or 0), rc['I%d' % r], rc['J%d' % r]))
        ws2 = wb['ServiceNow Daily Timecard']
        for ref, v in sorted(tc.items(), key=lambda kv: int(re.sub(r'\D', '', kv[0]))):
            print('timecard', ref, ws2['E%s' % re.sub(r'\D', '', ref)].value, '=', v)
        grand_tc = tc[max(tc, key=lambda k: int(re.sub(r'\D', '', k)))]
        grand_sn = rc['I%d' % max(int(re.sub(r'\D', '', k)) for k in rc if k.startswith('I'))]
        print('timecard grand total %.2f == reconciliation SN total %.2f ->' % (grand_tc, grand_sn), grand_tc == grand_sn)
        bad += grand_tc != grand_sn
    return bad


if __name__ == '__main__':
    paths = sys.argv[1:] or [os.path.join(HERE, f) for f in sorted(os.listdir(HERE)) if f.endswith('.xlsx')]
    sys.exit(1 if main(paths) else 0)
