"""Store each formula's computed result inside the workbook (the cached <v> value Excel keeps next to
a formula) so previewers that never recalculate still display totals. Excel and LibreOffice recompute
on open regardless (fullCalcOnLoad is set)."""
import os
import re
import shutil
import sys
import zipfile

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from verify_timecards import evaluate  # noqa: E402

CELL = re.compile(r'<c r="([A-Z]+\d+)"([^>]*)><f>([^<]*)</f>(?:<v>[^<]*</v>|<v\s*/>)?</c>')


def inject(path):
    wb = load_workbook(path)
    values = {}
    for ws in wb.worksheets:
        values[ws.title] = evaluate(ws)
    # worksheet part names in workbook order
    order = [ws.title for ws in wb.worksheets]
    tmp = path + '.tmp'
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            m = re.fullmatch(r'xl/worksheets/sheet(\d+)\.xml', item.filename)
            if m:
                vals = values[order[int(m.group(1)) - 1]]

                def repl(mm):
                    ref, attrs, formula = mm.group(1), mm.group(2), mm.group(3)
                    v = vals[ref]
                    if isinstance(v, str):
                        attrs = re.sub(r'\st="[^"]*"', '', attrs) + ' t="str"'
                        return '<c r="%s"%s><f>%s</f><v>%s</v></c>' % (ref, attrs, formula, v)
                    return '<c r="%s"%s><f>%s</f><v>%s</v></c>' % (ref, attrs, formula, ('%g' % v))

                text = data.decode('utf-8')
                text, n = CELL.subn(repl, text)
                data = text.encode('utf-8')
            dst.writestr(item, data)
    shutil.move(tmp, path)


if __name__ == '__main__':
    for p in sys.argv[1:]:
        inject(p)
        wb = load_workbook(p, data_only=True)
        for ws in wb.worksheets:
            cells = [(c.coordinate, c.value) for row in ws.iter_rows() for c in row if c.coordinate in evaluate(load_workbook(p)[ws.title])]
            print(os.path.basename(p), ws.title, 'cached', len(cells), 'e.g.', cells[:3], cells[-2:])
