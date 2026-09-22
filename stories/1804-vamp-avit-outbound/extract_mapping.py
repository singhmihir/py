"""Reads the sheet "SN to VAMP" of the client workbook and generates, without hand transcription:
vamp_mapping.json (the rows in sheet order, with the JSON structure of column H and the JSON field
name of column I) and "VAMP Field Check - Background Script.js", the read-only background script that
resolves every sheet row to the ServiceNow field behind it on the instance it runs on (same name,
then same label, then the u_ variant), checks the type and prints the one property value to use, as
for the CDP payloads: one servicenow_field=json_field pair per line, the ServiceNow field on the left
and the payload name of the sheet on the right as <column H>.<column I>. resolve_1804.py runs that
script on the PDI and writes properties.json."""
import os, json
import openpyxl
HERE = os.path.dirname(os.path.abspath(__file__))
ITEM_TABLE = 'sn_vul_app_vulnerable_item'
wb = openpyxl.load_workbook(os.path.join(HERE, 'SN to VAMP Mapping.xlsx'), data_only=True)
ws = wb['SN to VAMP']
header = [c for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
assert header[:5] == ['SN Fields', 'SN Payload ID', 'VAMPRequired?', 'SN Data Type', 'Table'], header
assert header[7:9] == ['JSON structure', 'JSON field name'], header
rows = []
for r in ws.iter_rows(min_row=2, values_only=True):
    if not any(c is not None for c in r):
        continue
    cells = [str(c).strip() if c is not None else '' for c in list(r) + [''] * (9 - len(r))]
    label, payload, required, dtype, table, notes, _carmack, structure, json_name = cells[:9]
    assert structure and json_name, cells
    rows.append({'table': table, 'structure': structure, 'json': json_name, 'payload': payload, 'label': label,
                 'type': dtype, 'required': required, 'notes': notes})
assert all(r['required'].lower() == 'yes' for r in rows), 'the generator assumes every sheet row is required'
# the sections of the payload, in the order the sheet introduces them; the JSON structure of column H
# is the payload name, lower case with underscores in place of spaces
sections, seen = [], set()
for r in rows:
    if r['table'] in seen:
        continue
    seen.add(r['table'])
    sections.append({'table': r['table'], 'json': r['structure'].strip().lower().replace(' ', '_')})
assert len({s['json'] for s in sections}) == len(sections), sections
for r in rows:
    same = [s for s in sections if s['table'] == r['table']]
    assert len(same) == 1 and same[0]['json'] == r['structure'].strip().lower().replace(' ', '_'), r
json.dump(rows, open(os.path.join(HERE, 'vamp_mapping.json'), 'w'), indent=1)
json.dump(sections, open(os.path.join(HERE, 'vamp_sections.json'), 'w'), indent=1)
check = '''// Checks the sheet "SN to VAMP" against this instance. For every sheet row it looks for the
// ServiceNow field behind the JSON field name: a field of that name, then a field with the sheet's
// label, then the u_ variant of the name. It checks the type against the sheet, prints one line per
// row, the one property value to use (the ServiceNow field on the left, the payload name on the
// right as <json structure>.<json field>) and the rows to raise. Read only; run as a background
// script in global scope.
var SHEET = %s;
var SECTIONS = %s;
var TYPES = { 'String': ['string'], 'Reference': ['reference'], 'Date/Time': ['glide_date_time'], 'Integer': ['integer'] };
var ITEM_TABLE = '%s';
var report = [], lines = [], raise = [], fieldLines = [], cache = {};

function fieldsOf(table) {
    if (!cache[table]) {
        var list = [], gr = new GlideRecord(table);
        if (gr.isValid()) {
            gr.initialize();
            var all = gr.getFields();
            for (var i = 0; i < all.size(); i++) {
                var ed = all.get(i).getED();
                list.push({ name: '' + ed.getName(), label: '' + ed.getLabel(), type: '' + ed.getInternalType(), reference: '' + (ed.getReference() || '') });
            }
        }
        cache[table] = list;
    }
    return cache[table];
}

function structureOf(table) {
    for (var s = 0; s < SECTIONS.length; s++)
        if (SECTIONS[s].table == table)
            return SECTIONS[s].json;
    return '';
}

function first(list, test) {
    for (var i = 0; i < list.length; i++)
        if (test(list[i]))
            return list[i];
    return null;
}

function resolve(row) {
    var fields = fieldsOf(row.table);
    var label = row.label.toLowerCase().replace('record number', 'number');
    var alt = row.json.indexOf('u_') == 0 ? row.json.substring(2) : 'u_' + row.json;
    var found = first(fields, function(f) { return f.name == row.json; });
    if (found)
        return { field: found, how: 'same name' };
    if (label)
        found = first(fields, function(f) { return f.label.toLowerCase() == label || f.label.toLowerCase() == row.label.toLowerCase(); });
    if (found)
        return { field: found, how: 'label "' + found.label + '"' };
    found = first(fields, function(f) { return f.name == alt; });
    if (found)
        return { field: found, how: 'name ' + alt };
    return null;
}

function candidates(row) {
    var words = row.label.toLowerCase().replace('record number', 'number').split(/\\s+/);
    var names = [];
    if (!row.label)
        return names;
    var fields = fieldsOf(row.table);
    for (var i = 0; i < fields.length; i++) {
        var l = fields[i].label.toLowerCase(), all = true;
        for (var w = 0; w < words.length; w++)
            if (l.indexOf(words[w]) < 0)
                all = false;
        if (all)
            names.push(fields[i].name + ' ("' + fields[i].label + '")');
    }
    return names;
}

for (var i = 0; i < SHEET.length; i++) {
    var row = SHEET[i], head = row.structure + '.' + row.json + ' (' + row.table + (row.label ? ', ' + row.label : '') + ', ' + row.type + '): ';
    var entry = { table: row.table, structure: row.structure, structure_json: structureOf(row.table), json: row.json, label: row.label, type: row.type, field: '', how: '', found_type: '', reference: '', type_ok: false, candidates: [] };
    if (!new GlideRecord(row.table).isValid()) {
        lines.push(head + 'TABLE MISSING');
        raise.push(row.table + ' does not exist');
    } else {
        var hit = resolve(row);
        if (!hit) {
            entry.candidates = candidates(row);
            lines.push(head + 'NOT FOUND - no field named ' + row.json + (row.label ? ', none labelled "' + row.label + '"' : '') + (entry.candidates.length ? '; similar labels: ' + entry.candidates.join(', ') : ''));
            raise.push(row.table + '.' + row.json + (row.label ? ' (' + row.label + ')' : '') + ' has no field on this instance');
        } else {
            entry.field = hit.field.name; entry.how = hit.how; entry.found_type = hit.field.type; entry.reference = hit.field.reference;
            entry.type_ok = (TYPES[row.type] || []).indexOf(hit.field.type) > -1;
            lines.push(head + hit.field.name + ' by ' + hit.how + ', type ' + hit.field.type + (hit.field.reference ? ' -> ' + hit.field.reference : '') + (entry.type_ok ? '' : ', TYPE DIFFERS FROM THE SHEET'));
            if (!entry.type_ok)
                raise.push(row.table + '.' + hit.field.name + ' is ' + hit.field.type + ' on this instance, the sheet says ' + row.type);
        }
    }
    report.push(entry);
}
for (var r = 0; r < report.length; r++) {
    var e = report[r];
    fieldLines.push((e.table == ITEM_TABLE ? '' : e.table + '.') + (e.field || e.json) + '=' + e.structure_json + '.' + e.json + ',');
}
gs.print('SN to VAMP field check on ' + gs.getProperty('instance_name') + ' - ' + SHEET.length + ' sheet rows\\n' + lines.join('\\n'));
gs.print('Property value to use for usem.vamp.fields.' + ITEM_TABLE + ' (ServiceNow field on the left, <json structure>.<json field> on the right, in sheet order; a row not found keeps the sheet name on the left until the field is known):\\n' + fieldLines.join('\\n'));
gs.print(raise.length ? 'To raise (' + raise.length + '):\\n- ' + raise.join('\\n- ') : 'Every sheet row resolves to a field of the sheet type.');
''' % (json.dumps([{'table': r['table'], 'structure': r['structure'], 'json': r['json'], 'label': r['label'], 'type': r['type']} for r in rows], indent=4),
       json.dumps(sections, indent=4), ITEM_TABLE)
open(os.path.join(HERE, 'VAMP Field Check - Background Script.js'), 'w').write(check)
print('sheet rows:', len(rows), '| sections:', [(s['table'], s['json']) for s in sections])
