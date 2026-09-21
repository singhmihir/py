"""Reads the tab "Outbound to CDP (consequence)" of the client workbook and generates, without hand
transcription: consequence_mapping.json (the rows marked CDP Required = Yes, in sheet order) and
"Consequence Field Check - Background Script.js", the read-only background script that resolves every
sheet row to the ServiceNow field behind it on the instance it runs on (the sheet's field name, then
the payload name, then the sheet's label, then the u_ variant, then the name without its u_ or x_
prefix) and prints the property value to use. resolve_1624.py runs that script on the PDI and writes
properties.json."""
import os, json
import openpyxl
HERE = os.path.dirname(os.path.abspath(__file__))
CONSEQUENCE = 'x_boar_bofa_usem_0_consequence'
wb = openpyxl.load_workbook(os.path.join(HERE, 'Terminology per Source Consolidation.xlsx'), data_only=True)
ws = wb['Outbound to CDP (consequence)']
header = [str(c).strip() if c is not None else '' for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
assert header[:13] == ['Field Label', 'Field Name', '', 'Kafka Topic', 'CDP Required?', 'New/Existing Field in CDP', 'Expected field type in CDP', 'JSON structure', 'JSON field name', '', 'Table', 'SN Field Label', 'SN Field Name'], header
rows, skipped = [], []
for r in ws.iter_rows(min_row=2, values_only=True):
    if not any(c is not None for c in r):
        continue
    cell = lambda i: str(r[i]).strip() if i < len(r) and r[i] is not None else ''
    label, field, topic, required, cdp_type, section, payload, table, sn_label, sn_field = cell(0), cell(1), cell(3), cell(4), cell(6), cell(7), cell(8), cell(10), cell(11), cell(12)
    if required.lower() != 'yes':
        if table or payload or label:
            skipped.append('%s.%s (%s)' % (table, payload, label))
        continue
    assert topic == 'sn_usem_consequence_outbound' and field == sn_field and label == sn_label and section in ('Consequence', 'Rule'), r[:13]
    rows.append({'table': table, 'payload': payload, 'label': label, 'field': field, 'type': cdp_type, 'section': section})
assert [r['table'] for r in rows] == sorted([r['table'] for r in rows], key=lambda t: t != CONSEQUENCE), 'consequence rows are expected before the rule rows'
json.dump(rows, open(os.path.join(HERE, 'consequence_mapping.json'), 'w'), indent=1)
check = '''// Checks the tab "Outbound to CDP (consequence)" against this instance. For every sheet row it looks
// for the ServiceNow field behind the payload name: the sheet's field name, then the payload name,
// then a field with the sheet's label, then the u_ variant, then the name without its u_ or x_
// prefix. It prints one line per row, the property value to use (ServiceNow field on the left,
// payload name on the right) and the rows to raise. Read only; run as a background script in global scope.
var SHEET = %s;
var CONSEQUENCE_TABLE = '%s';
var report = [], lines = [], raise = [], propertyLines = [], cache = {};

function fieldsOf(table) {
    if (!cache[table]) {
        var list = [], gr = new GlideRecord(table);
        if (gr.isValid()) {
            gr.initialize();
            var all = gr.getFields();
            for (var i = 0; i < all.size(); i++) {
                var ed = all.get(i).getED();
                list.push({ name: '' + ed.getName(), label: '' + ed.getLabel(), type: '' + ed.getInternalType(), reference: '' + (ed.getReference() || '') });
            }
            if (!first(list, function(f) { return f.name == 'sys_id'; }))
                list.push({ name: 'sys_id', label: 'Sys ID', type: 'GUID', reference: '' });
        }
        cache[table] = list;
    }
    return cache[table];
}

function first(list, test) {
    for (var i = 0; i < list.length; i++)
        if (test(list[i]))
            return list[i];
    return null;
}

function resolve(row) {
    var fields = fieldsOf(row.table);
    var label = row.label.toLowerCase();
    var bare = row.field.replace(/^(u_|x_)/, '');
    var attempts = [
        ['same name', function(f) { return f.name == row.field; }],
        ['payload name', function(f) { return f.name == row.payload; }],
        ['label "' + row.label + '"', function(f) { return f.label.toLowerCase() == label; }],
        ['name u_' + bare, function(f) { return f.name == 'u_' + bare; }],
        ['name ' + bare, function(f) { return f.name == bare; }]
    ];
    for (var i = 0; i < attempts.length; i++) {
        var found = first(fields, attempts[i][1]);
        if (found)
            return { field: found, how: attempts[i][0] };
    }
    return null;
}

function candidates(row) {
    var words = row.label.toLowerCase().split(/\\s+/);
    var names = [];
    var fields = fieldsOf(row.table);
    for (var i = 0; i < fields.length; i++) {
        var l = fields[i].label.toLowerCase(), all = true;
        for (var w = 0; w < words.length; w++)
            if (l.indexOf(words[w]) < 0)
                all = false;
        if (all)
            names.push(fields[i].name + ' ("' + fields[i].label + '")');
    }
    return names;
}

for (var i = 0; i < SHEET.length; i++) {
    var row = SHEET[i], head = row.table + '.' + row.field + ' -> ' + row.payload + ' (' + row.label + '): ';
    var entry = { table: row.table, payload: row.payload, label: row.label, sheet_field: row.field, field: '', how: '', found_type: '', reference: '', candidates: [] };
    if (!new GlideRecord(row.table).isValid()) {
        lines.push(head + 'TABLE MISSING');
        raise.push(row.table + ' does not exist');
    } else {
        var hit = resolve(row);
        if (!hit) {
            entry.candidates = candidates(row);
            lines.push(head + 'NOT FOUND - no field named ' + row.field + ' or ' + row.payload + ', none labelled "' + row.label + '"' + (entry.candidates.length ? '; similar labels: ' + entry.candidates.join(', ') : ''));
            raise.push(row.table + '.' + row.field + ' (' + row.label + ') has no field on this instance');
        } else {
            entry.field = hit.field.name; entry.how = hit.how; entry.found_type = hit.field.type; entry.reference = hit.field.reference;
            lines.push(head + hit.field.name + ' by ' + hit.how + ', type ' + hit.field.type + (hit.field.reference ? ' -> ' + hit.field.reference : ''));
        }
    }
    report.push(entry);
}
for (var p = 0; p < 2; p++)
    for (var r = 0; r < report.length; r++) {
        var e = report[r];
        if ((p == 0) != (e.table == CONSEQUENCE_TABLE))
            continue;
        propertyLines.push((e.table == CONSEQUENCE_TABLE ? '' : e.table + '.') + (e.field || e.sheet_field) + '=' + e.payload + ',');
    }
gs.print('Outbound to CDP (consequence) field check on ' + gs.getProperty('instance_name') + ' - ' + SHEET.length + ' sheet rows\\n' + lines.join('\\n'));
gs.print('Property value to use for usem.consequence.fields.' + CONSEQUENCE_TABLE + ' (a row not found keeps the sheet name until the field is known):\\n' + propertyLines.join('\\n'));
gs.print(raise.length ? 'To raise (' + raise.length + '):\\n- ' + raise.join('\\n- ') : 'Every sheet row resolves to a field on this instance.');
''' % (json.dumps([{'table': r['table'], 'payload': r['payload'], 'label': r['label'], 'field': r['field']} for r in rows], indent=4), CONSEQUENCE)
open(os.path.join(HERE, 'Consequence Field Check - Background Script.js'), 'w').write(check)
print('sheet rows required:', len(rows), '| tables:', sorted(set(r['table'] for r in rows)), '| skipped (not required):', skipped)
