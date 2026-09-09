"""Builds 'Qualys CI Lookup Rules - Code Line Explanations.xlsx' from rules/*.js:
an Overview sheet (one row per rule: purpose, input, returns, place in the chain) and
one sheet per rule listing every code line with its explanation, grouped by stage."""
import glob, json, os, re
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, 'Qualys CI Lookup Rules - Code Line Explanations.xlsx')
TEXT = Font(name='Arial', size=10)
BOLD = Font(name='Arial', size=10, bold=True)
CODE = Font(name='Courier New', size=9)
HEAD_FILL = PatternFill('solid', fgColor='1F3864'); HEAD_FONT = Font(name='Arial', size=10, bold=True, color='FFFFFF')
STAGE_FILL = PatternFill('solid', fgColor='DDEBF7')
THIN = Side(style='thin', color='BFBFBF'); BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='top')


def parse(path):
    text = open(path).read()
    head, body = text.split('*/', 1)
    order = os.path.basename(path).split('_', 1)[0]
    hl = [l[3:] if l.startswith('   ') else l for l in head.split('\n')]
    name = hl[0][3:].strip()
    # header paragraphs: purpose (until the first blank line), then labelled lines
    i = 2; purpose = []
    while i < len(hl) and hl[i].strip():
        purpose.append(hl[i].strip()); i += 1
    labelled = {}; cur = None
    for l in hl[i:]:
        m = re.match(r'(Input|Returns|Before|Reaches|After)\s*: (.*)', l)
        if m:
            cur = m.group(1); labelled[cur] = m.group(2).strip()
        elif cur and l.startswith('         ') and l.strip():
            labelled[cur] += ' ' + l.strip()
        elif l.startswith('Place in the chain'):
            cur = None
    field = re.search(r'sourceValue is the (\S+) field', labelled.get('Input', '')).group(1)
    # body: stage notes, plain comments, code lines
    lines = body.split('\n')
    rows = []; pending = []; i = 0
    while i < len(lines):
        l = lines[i]; s = l.strip()
        m = re.match(r'// -- (.*?) -+$', s)
        if m:
            title = m.group(1); i += 1; note = []
            while i < len(lines) and lines[i].strip().startswith('//') and not re.match(r'// -- ', lines[i].strip()):
                note.append(lines[i].strip()[2:].strip()); i += 1
            rows.append(('stage', title, ' '.join(note)))
            pending = []; continue
        if s.startswith('//'):
            pending.append(s[2:].strip()); i += 1; continue
        if s == '':
            i += 1; continue
        code, trail = l, ''
        tm = re.search(r'\s*// (.*)$', l)
        if tm: code, trail = l[:tm.start()], tm.group(1)
        expl = ' '.join(pending)
        if trail: expl = (expl + '\n' if expl else '') + trail
        cs = code.strip()
        if rows and rows[-1][0] == 'code' and continues(rows[-1][1], cs):
            k, c, e = rows[-1]; rows[-1] = (k, c + '\n' + code.rstrip(), (e + '\n' + expl) if (e and expl) else (e or expl))
        else:
            rows.append(('code', code.rstrip(), expl))
        pending = []; i += 1
    out = []
    for k, c, e in rows:
        if k == 'code' and not e:
            e = default_expl(c)
        out.append((k, c, e))
    return dict(order=order, name=name, purpose=' '.join(purpose), field=field, input=labelled.get('Input', ''), returns=labelled.get('Returns', ''),
                before=labelled.get('Before', ''), reaches=labelled.get('Reaches', ''), after=labelled.get('After', ''), rows=out)

CONTROL = re.compile(r'^(if|for|while)\s*\(.*\)$')


def continues(prev_code, cur):
    """True when the physical line cur belongs to the statement started in prev_code."""
    last = prev_code.split('\n')[-1].strip()
    if last.startswith('(function') or cur.startswith('})('):
        return False
    if CONTROL.match(last) and not cur.startswith('}'):
        return True                                   # single-line body of an if/for/while
    if not (last.endswith(';') or last.endswith('{') or last.endswith('}') or CONTROL.match(last)):
        return True                                   # statement not finished (operators, lists)
    return False


DEFAULTS = [
    (r'^\(function process', 'Entry point. The framework calls process(rule, sourceValue, sourcePayload) once per scanned host and uses what it returns: a CI sys_id links the host, null hands the host to the next rule.'),
    (r'^\}\)\(rule', 'Runs the function immediately with the values the framework supplied.'),
    (r'^function classFor', 'Helper: translates the OS text into the CMDB class the CI lives in (see the examples above).'),
    (r"if \(!os\) return ''", 'No OS text -> no class.'),
    (r"var s = \('' \+ os\)\.toLowerCase\(\)", 'Lower-case the OS text so that the key-word checks are case-insensitive.'),
    (r"return 'cmdb_ci_esx_server'", 'OS text contains "esx" (for example "VMware ESXi 7.0.3 build 24723872") -> ESX Server.'),
    (r"'cmdb_ci_win_server' : 'cmdb_ci_computer'", 'OS text contains "windows": with "server" -> Windows Server, otherwise -> Computer (a workstation).'),
    (r"return 'cmdb_ci_aix_server'", 'OS text contains "aix" (for example "AIX 7.3 TL3") -> AIX Server.'),
    (r"return 'cmdb_ci_solaris_server'", 'OS text contains "solaris" or "sunos" -> Solaris Server.'),
    (r"return 'cmdb_ci_hpux_server'", 'OS text contains "hp-ux" -> HPUX Server.'),
    (r"return 'cmdb_ci_storage_server'", 'OS text contains "netapp" or "ontap" -> Storage Server.'),
    (r"return 'cmdb_ci_printer'", 'OS text contains "printer", "laserjet" or "jetdirect" -> Printer.'),
    (r"return 'cmdb_ci_linux_server'", 'OS text contains one of the Linux family words ("red hat", "linux", "centos", "ubuntu", "suse", "debian", "fedora", "euleros", "oracle enterprise", "amazon") -> Linux Server.'),
    (r"return 'cmdb_ci_netgear'", 'OS text contains "nx-os", "catos" or "cisco" -> Network Gear.'),
    (r"^return '';$", 'Nothing matched -> no class preference.'),
    (r'^function isLoadBalancer', 'Helper: true when the CI with this sys_id is a Load Balancer.'),
    (r"var lb = new GlideRecord\('cmdb_ci_lb'\)", 'Search on the Load Balancer class.'),
    (r'return lb\.isValid\(\) && lb\.get\(id\)', 'True only when the class exists and a load balancer with this sys_id is found.'),
    (r'^function pick(Fqdn|Combo)\(table\)', 'Helper: the search-and-decide logic described in the stage banner, run on the table passed in.'),
    (r'^var gr = new GlideRecord\(table\)', 'Search on the table passed in (the class and every class beneath it).'),
    (r'^if \(!gr\.isValid\(\)\)', 'The class is not installed on this instance -> decline rather than fail.'),
    (r"^if \(ignore\)\n\s*\w+\.addQuery\('sys_class_name', 'NOT IN', ignore\)", 'When the ignore list is filled, exclude those classes from the search.'),
    (r"addQuery\('sys_class_name', 'NOT IN', ignore\)", 'When the ignore list is filled, exclude those classes from the search.'),
    (r"addQuery\('cmdb_ci\.sys_class_name', 'NOT IN', ignore\)", 'When the ignore list is filled, exclude adapters whose owning CI is in an ignored class.'),
    (r"addQuery\('nic\.cmdb_ci\.sys_class_name', 'NOT IN', ignore\)", 'When the ignore list is filled, exclude records whose owning CI is in an ignored class.'),
    (r"addQuery\('ip_address\.nic\.cmdb_ci\.sys_class_name', 'NOT IN', ignore\)", 'When the ignore list is filled, exclude chains that end on an ignored class.'),
    (r'^\w+\.query\(\);$', 'Runs the search.'),
    (r"^gr\.addQuery\('serial_number', serial\)", 'Condition: serial_number equals the cleaned serial (exact match).'),
    (r"^gr\.addQuery\('name', host\)", 'Condition: name equals the short hostname (CMDB names compare case-insensitively).'),
    (r"^gr\.addQuery\('name', fqdn\)", 'Condition: name equals the complete DNS name.'),
    (r"^gr\.addQuery\('ip_address', ip\)", 'Condition: ip_address equals the scanned address.'),
    (r"^gr\.addQuery\('fqdn', fqdn\)", 'Condition: fqdn equals the scanned DNS name.'),
    (r"^nic\.addQuery\('ip_address', ip\)", 'Condition: the adapter carries the scanned address.'),
    (r"^ipGr\.addQuery\('ip_address', ip\)", 'Condition: the IP Address record holds the scanned address.'),
    (r"^nic\.addQuery\('mac_address', 'IN'", 'Condition: the adapter MAC is one of the four spellings.'),
    (r"^ph\.addQuery\('mac_address', 'IN'", 'Condition: the phone record MAC is one of the four spellings.'),
    (r'^while \(\w+\.next\(\)\) \{', 'Loop over every row of the result.'),
    (r'^\}$', 'End of the block.'),
    (r"^var suffixes = \['ilo'", 'The label suffixes that mark a management controller.'),
    (r"^var markers = \['ilo'", 'The words that name a controller in the OS text.'),
    (r"^var domains = \['\.network\.'\]", 'The domains under which network devices are scanned per interface.'),
    (r"^var markers = \['vlan'", 'The label segments that mark an interface or VLAN address.'),
    (r"^var osMarkers = \['f5'", 'The load balancer products looked for in the OS text.'),
    (r"^var labelMarkers = \['vip'", 'The label segments that mark a virtual IP.'),
    (r'^\} else \{$', 'Otherwise.'),
    (r'^var pref = classFor\(sourcePayload\.OS\)', 'The class the scanned OS points at; empty when the OS is unknown or a multi-guess fingerprint.'),
    (r'^if \(!pref\)\n\s*return null', 'No class from the OS -> this class-scoped rule declines and the hardware-wide rule takes over.'),
    (r'^if \(\w+\.hasNext\(\)\)\n\s*return null', 'A second row means the value is shared by two CIs -> decline rather than guess.'),
    (r'^if \(!\w+\.next\(\)\)\n\s*return null', 'No row at all -> decline.'),
    (r'^var match = \w+\.getUniqueValue\(\)', 'The sys_id of the CI found.'),
    (r'^return match;', 'Exactly one CI -> return its sys_id; the framework links the host to it and stops the chain.'),
    (r"^var cls = '' \+ \w+\.getValue\('sys_class_name'\)", 'The class of the CI found, checked later against the scanned OS.'),
    (r'^var id = \w+\.getUniqueValue\(\)', 'The sys_id of the CI found.'),
    (r"^var parts = \('' \+ gs\.getProperty\(name, fallback\)\)\.split\(','\)", 'Read the property (or the default shipped with the rule) and split it on commas.'),
    (r'^var out = \[\];$', 'The cleaned list.'),
    (r'^for \(var i = 0; i < parts\.length; i\+\+\) \{', 'Loop over the comma separated entries.'),
    (r'^var item = parts\[i\]\.trim\(\)\.toLowerCase\(\)', 'Trim and lower-case one entry.'),
    (r'^if \(item\)\n\s*out\.push\(item\)', 'Keep non-empty entries.'),
    (r'^return out;', 'The list of lower-cased entries.'),
    (r"^var dot = full\.indexOf\('\.'\)", 'Position of the first dot, which separates the short hostname from the domain.'),
    (r"^gr\.addQuery\('name', base\)", 'Condition: name equals the derived name.'),
    (r'^for \(var i = 0; i < words\.length; i\+\+\) \{', 'Loop over the listed marker words.'),
    (r'^var w = words\[i\];', 'One marker word.'),
    (r'^if \(segment == w\)\n\s*return true', 'The segment is the word itself ("vlan").'),
    (r'^if \(segment\.indexOf\(w\) == 0 && ', 'The segment is the word followed by digits only ("vlan705", "v201").'),
    (r'^if \(w\.length >= 3 && segment\.length > w\.length', 'For words of three letters or more, the segment ends in the word ("multihostvip").'),
    (r'^return false;', 'No marker word fits this segment.'),
    (r'^var evidence = false;', 'Set to true as soon as one piece of evidence is found.'),
    (r'^if \(!evidence\)\n\s*return null', 'No evidence -> this host is not one the rule is for; decline untouched.'),
    (r'^var m = label\.match\(/\^sep', 'Strict pattern: "sep" plus exactly twelve hexadecimal characters; anything else is not a phone label.'),
    (r'^if \(!ph\.hasNext\(\)\)\n\s*return byMac', 'Exactly one phone carries the MAC on its own record -> match.'),
    (r'^if \(byName\.hasNext\(\)\)\n\s*return null', 'Two phones with one device name -> never guess.'),
    (r"^gr\.addQuery\('dns_name\.name', fqdn\)", 'Condition: the DNS Name record at the start of the chain is named with the scanned name.'),
    (r"^var dash = label\.lastIndexOf\('-'\)", 'Position of the last hyphen; the controller suffix sits after it.'),
    (r"^var os = \('' \+ \(sourcePayload\.OS \|\| ''\)\)\.toLowerCase\(\)", 'The scanned OS text, lower-cased; empty when Qualys reported none.'),
    (r"^var suffixes = list\('usem\.ci_lookup\.mgmt_suffixes'", 'Controller suffixes, from the property (defaults shipped with the rule).'),
    (r"^var markers = list\('usem\.ci_lookup\.mgmt_os_markers'", 'Controller words looked for in the OS text, from the property.'),
    (r"^var base = '';", 'The server name once derived; stays empty without controller evidence.'),
    (r'^else\n\s*for \(var i = 0; i < markers\.length; i\+\+\)', 'No listed suffix: when the OS text names a controller, strip the last segment whatever it is.'),
    (r'^break;', 'One controller word is enough.'),
    (r'^if \(!base\)\n\s*return null', 'No controller evidence -> decline untouched.'),
    (r'^if \(lb\.isValid\(\) && lb\.get\(match\)\)\n\s*return null', 'The CI found is a load balancer -> refuse it.'),
    (r"^var label = full\.split\('\.'\)\[0\]", 'The first label of the DNS name.'),
    (r"^var label = dns\.split\('\.'\)\[0\]", 'The first label of the DNS name, empty when there is no DNS name.'),
    (r"^var domains = list\('usem\.ci_lookup\.interface_domains'", 'Interface domains, from the property.'),
    (r"^var markers = list\('usem\.ci_lookup\.interface_markers'", 'Interface marker words, from the property.'),
    (r'^for \(var d = 0; d < domains\.length; d\+\+\)', 'Evidence when the DNS name contains a listed interface domain.'),
    (r'^for \(var s = 1; s < segments\.length; s\+\+\)', 'Evidence when a segment after the first is a listed marker.'),
    (r"^var tables = \['cmdb_ci_netgear', 'cmdb_ci_lb'\]", 'The two branches that hold network devices.'),
    (r'^for \(var k = segments\.length - 1; k >= 1; k--\) \{', 'From the longest prefix (all segments but the last) down to the first segment alone.'),
    (r'^var hits = \[\];', 'Devices named exactly like this prefix.'),
    (r'^for \(var t = 0; t < tables\.length; t\+\+\) \{', 'Search both branches.'),
    (r'^var gr = new GlideRecord\(tables\[t\]\)', 'Search on one branch and every class beneath it.'),
    (r'^while \(gr\.next\(\) && hits\.length < 2\)', 'Collect up to two hits; two is already one too many.'),
    (r'^return hits\[0\];', 'Exactly one device carries this prefix -> match.'),
    (r"^var osMarkers = list\('usem\.ci_lookup\.vip_os_markers'", 'Load balancer product words looked for in the OS text, from the property.'),
    (r"^var labelMarkers = list\('usem\.ci_lookup\.vip_markers'", 'VIP marker words for the label segments, from the property.'),
    (r'^for \(var m = 0; m < osMarkers\.length; m\+\+\)', 'Evidence when the OS text contains a listed load balancer word.'),
    (r"^var segments = label \? label\.split\('-'\) : \[\]", 'The hyphen segments of the label; none without a DNS name.'),
    (r'^for \(var s = 0; s < segments\.length; s\+\+\)', 'Evidence when any segment is a listed VIP marker.'),
    (r'^function one\(field, value\)', 'Helper: the one Load Balancer Service whose field equals the value; undefined when nothing to search or nothing found, null when two are found.'),
    (r"^var gr = new GlideRecord\('cmdb_ci_lb_service'\)", 'Search on the Load Balancer Service class and every class beneath it.'),
    (r'^gr\.addQuery\(field, value\)', 'Condition: the field named by the step equals the value.'),
    (r"^var steps = \[\['fqdn', dns\]", 'The evidence in order of strength: the DNS name as fqdn, as name, its label as name, then the address.'),
    (r'^for \(var t = 0; t < steps\.length; t\+\+\) \{', 'Try the steps in order.'),
    (r'^var found = one\(steps\[t\]\[0\], steps\[t\]\[1\]\)', 'Run one step.'),
    (r'^if \(found === null\)\n\s*return null', 'Two services at this step -> decline; a weaker step must not override an ambiguous stronger one.'),
    (r'^if \(found\)\n\s*return found', 'Exactly one service -> match.'),
    (r'^ids\.push\(gr\.getUniqueValue\(\)\)', 'Remember the sys_id of this CI.'),
    (r"^if \(ip && gr\.getValue\('ip_address'\) == ip\)", 'When this CI also carries the scanned IP, remember it as an IP-confirmed candidate.'),
    (r'^good\.push\(gr\.getUniqueValue\(\)\)', 'The domain evidence agrees -> remember the sys_id of this CI.'),
    (r'^if \(cifqdn == full', 'Keep the CI only when its own domain information confirms the scanned domain (one of the three tests).'),
    (r'^var phones = \{\}, count = 0, first = null', 'Bookkeeping: the phones already counted, how many, and the first one.'),
    (r'^var owners = \{\}', 'Bookkeeping: the owning CIs already counted.'),
    (r'^var count = 0, first = null', 'How many distinct owners were seen, and the first one.'),
    (r"^var owner = nic\.getValue\('cmdb_ci'\)", 'The CI that owns this adapter.'),
    (r"^var phone = new GlideRecord\('cmdb_ci_ip_phone'\)", 'Prepare a check that the owner is an IP Phone.'),
    (r'^phones\[owner\] = true;', 'Mark this phone as counted.'),
    (r'^owners\[owner\] = true;', 'Mark this CI as counted.'),
    (r'^count\+\+;', 'One more distinct owner.'),
    (r'^if \(count == 1\)\n\s*first = owner', 'Remember the first owner seen.'),
    (r'^if \(count == 1\)\n\s*return first', 'Exactly one owner -> that CI is the match.'),
    (r'^if \(count == 1 && !isLoadBalancer\(first\)\)', 'Exactly one owner and it is not a load balancer -> that CI is the match.'),
    (r'^if \(count > 1\) \{', 'Several owners: try the IP tie-break.'),
    (r'^var confirmed = Object\.keys\(ipOwners\)', 'The owners reached through the scanned IP.'),
    (r'^if \(confirmed\.length == 1\)', 'Exactly one owner confirmed by the address -> that CI is the match.'),
    (r'^ipOwners\[owner\] = true', 'Note that this owner is reached through the scanned IP.'),
    (r'^var byMac = ph\.getUniqueValue\(\)', 'The sys_id of the phone found by MAC.'),
    (r'^if \(ph\.next\(\)\) \{', 'At least one phone carries the MAC on its record.'),
    (r'^if \(!byName\.next\(\)\)', 'No phone with that device name -> decline.'),
    (r'^var named = byName\.getUniqueValue\(\)', 'The sys_id of the phone found by name.'),
    (r'^return named;', 'Exactly one phone -> its sys_id goes back to the framework.'),
    (r'^return null;$', 'Nothing usable -> decline; the next rule gets its turn.'),
    (r'^if \(!gr\.next\(\)\)', 'Empty result -> decline.'),
    (r'^var id = gr\.getUniqueValue\(\)', 'The sys_id of the CI found.'),
    (r'^if \(pref\) \{', 'Only when the OS gave a class preference:'),
    (r'^var chk = new GlideRecord\(pref\)', 'Search on the class the OS implies.'),
    (r'^if \(!\(chk\.isValid\(\) && chk\.get\(id\)\) && !generic\[cls\]\)', 'The CI is neither in the OS-implied class nor generically classed -> contradiction -> decline.'),
    (r'^return id;', 'The CI passed every check -> its sys_id goes back to the framework.'),
    (r'^if \(!m\)', 'The label is not a phone label -> decline.'),
    (r'^var candidates = \[', 'The four spellings of the MAC address.'),
    (r'^var pairs = \[\];', 'Will hold the six two-character pairs of the MAC.'),
    (r'^if \(!host\)', 'Empty hostname (a name starting with a dot) -> decline.'),
    (r'^\}$', ''),
]


EXTRA = [
    (r'^if \(!pref\)', 'The OS gave no class -> decline so that the hardware-wide rule gets its turn.'),
    (r'^if \(serial\.length < 4', 'Too short or a placeholder serial -> decline.'),
    (r"^var ph = new GlideRecord\('cmdb_ci_ip_phone'\)", 'Search on the IP Phone class.'),
    (r"^var byName = new GlideRecord\('cmdb_ci_ip_phone'\)", 'Search on the IP Phone class.'),
    (r"^var gr = new GlideRecord\('cmdb_ip_address_dns_name'\)", 'Search on the IP Address to DNS Name link table.'),
    (r"^var nic = new GlideRecord\('cmdb_ci_network_adapter'\)", 'Search on the Network Adapter table.'),
    (r"^var ipGr = new GlideRecord\('cmdb_ci_ip_address'\)", 'Search on the IP Address table.'),
    (r"^byName\.addQuery\('name', label\.toUpperCase\(\)\)", 'Condition: the phone is named with the SEP device name.'),
    (r'^var hex = m\[1\]', 'The twelve hexadecimal characters captured after "sep".'),
    (r'^for \(var i = 0; i < 12; i \+= 2\)', 'Cut the twelve characters into six pairs.'),
    (r"^var colon = pairs\.join\(':'\)", 'Join the pairs with colons.'),
    (r'^var generic = \{', 'The generically classed CI classes (see the comment above).'),
    (r'^if \(isLoadBalancer\(id\)\)', 'The single owner is a load balancer -> decline.'),
    (r'^var match = gr\.getUniqueValue\(\)', 'The sys_id of the CI found.'),
    (r'^return match;', 'Exactly one CI -> its sys_id goes back to the framework.'),
    (r'^return pick(Fqdn|Combo)\(', 'Run the helper on the chosen table and hand its answer back to the framework.'),
    (r'^var cifqdn = ', "The CI's own fqdn, lower-cased."),
    (r'^var cidom = ', "The CI's own dns_domain, lower-cased."),
    (r'^if \(ids\.length == 1\)', 'One CI carries the FQDN -> match.'),
    (r'^if \(good\.length == 1\)', 'One CI confirmed by name and domain -> match.'),
    (r'^if \(ids\.length > 1 && ipHits\.length == 1\)', 'Several CIs, exactly one confirmed by the scanned IP -> that one.'),
    (r'^if \(good\.length > 1 && ipHits\.length == 1\)', 'Several confirmed CIs, exactly one also carries the scanned IP -> that one.'),
    (r'^ipHits\.push\(gr\.getUniqueValue\(\)\)', 'Remember it as an IP-confirmed candidate.'),
    (r'^var ids = \[\];', 'Will hold every CI carrying the FQDN.'),
    (r'^var good = \[\];', 'Will hold the CIs whose domain evidence agrees.'),
    (r'^var ipHits = \[\];', 'Will hold those of them that also carry the scanned IP.'),
    (r'^var owners = \{\}, ipOwners = \{\};', 'Bookkeeping: the owning CIs already counted, and those reached through the scanned IP.'),
    (r"^var owner = '' \+ ", 'The CI at the end of this chain.'),
    (r'^if \(!owners\[owner\]\) \{', 'First time this CI is seen:'),
    (r"^nic\.addNotNullQuery\('cmdb_ci'\)", 'Condition: the adapter must belong to a CI.'),
    (r"^ipGr\.addNotNullQuery\('nic\.cmdb_ci'\)", 'Condition: the chain must end on a CI.'),
    (r"^gr\.addNotNullQuery\('ip_address\.nic\.cmdb_ci'\)", 'Condition: the chain must end on a CI.'),
    (r'^if \(phone\.get\(owner\) && !phones\[owner\]\) \{', 'Only IP phones count, and each phone only once.'),
    (r'^if \(count == 1\)$', 'Exactly one owner:'),
]
DEFAULTS[-1:-1] = EXTRA


def default_expl(code):
    cs = '\n'.join(l.strip() for l in code.split('\n'))
    for rx, e in DEFAULTS:
        if re.search(rx, cs):
            return e
    return ''


def header_row(ws, cols, widths):
    for j, (h, w) in enumerate(zip(cols, widths), 1):
        c = ws.cell(row=1, column=j, value=h); c.font = HEAD_FONT; c.fill = HEAD_FILL; c.alignment = WRAP; c.border = BORDER
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = 'A2'


rules = [parse(f) for f in sorted(glob.glob(os.path.join(HERE, 'rules', '*.js')))]
wb = Workbook(); ov = wb.active; ov.title = 'Overview'
header_row(ov, ['Order', 'Rule', 'Source field', 'Purpose', 'Input', 'Returns', 'Runs before it', 'Which hosts reach it', 'Runs after it'],
           [8, 30, 14, 55, 40, 50, 40, 40, 40])
for r, x in enumerate(rules, 2):
    vals = [int(x['order']), x['name'], x['field'], x['purpose'], x['input'], x['returns'], x['before'], x['reaches'], x['after']]
    for j, v in enumerate(vals, 1):
        c = ov.cell(row=r, column=j, value=v); c.font = TEXT; c.alignment = WRAP; c.border = BORDER
ov.auto_filter.ref = 'A1:I%d' % (len(rules) + 1)
for x in rules:
    ws = wb.create_sheet(('%s %s' % (x['order'], x['name'].replace('USEM ', '')))[:31])
    header_row(ws, ['#', 'Stage', 'Code', 'Explanation'], [5, 34, 70, 90])
    r = 2; n = 0
    for kind, code, expl in x['rows']:
        if kind == 'stage':
            for j, v in enumerate(['', code, '', expl], 1):
                c = ws.cell(row=r, column=j, value=v); c.font = BOLD if j == 2 else TEXT; c.fill = STAGE_FILL; c.alignment = WRAP; c.border = BORDER
        else:
            n += 1
            for j, v in enumerate([n, '', code, expl], 1):
                c = ws.cell(row=r, column=j, value=v); c.font = CODE if j == 3 else TEXT; c.alignment = WRAP; c.border = BORDER
        r += 1
wb.save(OUT)
print('workbook written:', OUT, '| sheets', len(wb.sheetnames), '| code rows', sum(1 for x in rules for k, _, _ in x['rows'] if k == 'code'))
