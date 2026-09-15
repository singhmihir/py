"""Builds 'Qualys Unmatched Hosts - Analysis.xlsx' from unmatched_profile.json:
the unmatched population by the evidence each host carries, which rule (existing
or new) can resolve each group once the CMDB holds the CI, and what remains a
CMDB data matter. Sample percentages are applied to the full population count."""
import json, os, re
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
HERE = os.path.dirname(os.path.abspath(__file__))
P = json.load(open(os.path.join(HERE, 'unmatched_profile.json'))); S = P['sample']; N = len(S)
TOTAL = 281700
TEXT = Font(name='Arial', size=10); BOLD = Font(name='Arial', size=10, bold=True)
HEAD_FILL = PatternFill('solid', fgColor='1F3864'); HEAD_FONT = Font(name='Arial', size=10, bold=True, color='FFFFFF')
THIN = Side(style='thin', color='BFBFBF'); BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN); WRAP = Alignment(wrap_text=True, vertical='top')
def label(r): return r['dns'].split('.')[0].lower() if r['dns'] else ''
def segs(r): return label(r).split('-') if label(r) else []
def is_marker(seg, words):
    for w in words:
        if seg == w or (seg.startswith(w) and seg[len(w):].isdigit()) or (len(w) >= 3 and len(seg) > len(w) and seg.endswith(w)): return True
    return False
MGMT = 'ilo,ilom,idrac,drac,ipmi,bmc,oob,mgmt,imm,cimc,rmm,con'.split(','); MGMT_OS = 'ilo,ilom,idrac,drac,remote access controller,imm,cimc,bmc,ipmi,lights out'.split(',')
IFM = 'vlan,v,hsrp,vrrp,po,eth,gi,te,lo,mgmt,aom,vs,fab'.split(','); VIPM = ['vip', 'vs']; VIPOS = ['f5', 'big-ip', 'big ip', 'netscaler']
def cat(r):
    lab, os_ = label(r), r['os'].lower(); sg = segs(r)
    if not r['dns']: return 'IP only (no DNS name)'
    if re.match(r'^sep[0-9a-f]{12}$', lab): return 'Cisco IP phone with SEP label'
    if 'ip phone' in os_: return 'Cisco IP phone, other label'
    if any(m in os_ for m in VIPOS) or any(is_marker(s, VIPM) for s in sg): return 'Virtual IP of a load balancer'
    if (len(sg) > 1 and sg[-1] in MGMT) or (any(m in os_ for m in MGMT_OS) and len(sg) > 1): return 'Management controller (iLO, iDRAC, ILOM)'
    if '.network.' in r['dns'].lower() or any(is_marker(s, IFM) for s in sg[1:]): return 'Network device interface or VLAN address'
    if any(k in os_ for k in ('printer', 'laserjet', 'lexmark', 'jetdirect')): return 'Printer'
    if any(k in os_ for k in ('camera', 'crestron', 'qnx', 'pdu', 'avocent', 'netbotz', 'embedded')): return 'Appliance or embedded device'
    if any(k in os_ for k in ('nexus', 'ios', 'cisco device', 'catos', 'nx-os', 'juniper', 'arista')): return 'Network device by its own name'
    if 'windows 1' in os_ or 'windows 10' in os_ or 'windows 11' in os_: return 'Windows workstation'
    if 'netapp' in os_ or 'ontap' in os_: return 'Storage system'
    return 'Server or other host with a plain DNS name'
RULES = {'IP only (no DNS name)': ('700 to 740 (IP address rules), 460 for VIPs', 'Yes, when the CMDB holds the address on the CI, an adapter or an IP Address record'),
         'Cisco IP phone with SEP label': ('200', 'Yes, when the CMDB holds IP Phone CIs with the MAC or the SEP device name'),
         'Cisco IP phone, other label': ('none', 'No: the label is not a SEP name and carries no attribute a phone CI would hold'),
         'Virtual IP of a load balancer': ('460 (new)', 'Yes, when the CMDB holds Load Balancer Service CIs with the fqdn, name or address'),
         'Management controller (iLO, iDRAC, ILOM)': ('420 (new)', 'Yes, when the server CI is named after the label without its suffix'),
         'Network device interface or VLAN address': ('430 (new)', 'Yes, when the device CI is named after the leading part of the label'),
         'Printer': ('250 to 450 by name if a Printer CI exists', 'Only when the CMDB holds printers'),
         'Appliance or embedded device': ('250 to 450 by name if a CI exists', 'Only when the CMDB holds these devices'),
         'Network device by its own name': ('250 to 450 by name, 700 to 740 by address', 'Yes, when the CMDB holds the device under that name'),
         'Windows workstation': ('400/410 by short name', 'Only when the CMDB holds workstations'),
         'Storage system': ('250 to 450 by name', 'Yes, when the CMDB holds the storage CI under that name'),
         'Server or other host with a plain DNS name': ('250 to 450 by name, 700 to 740 by address', 'Yes, when a CI with that name, fqdn or address exists; a host in this group is unmatched because the CMDB has no such CI or has two')}
cats = Counter(cat(r) for r in S)
wb = Workbook(); ws = wb.active; ws.title = 'Population by evidence'
heads = ['Group', 'Share of sample', 'Estimated hosts (of %s)' % format(TOTAL, ','), 'Rule that resolves the group', 'Can a rule resolve it?', 'Example DNS name', 'Example OS text']
for j, h in enumerate(heads, 1):
    c = ws.cell(row=1, column=j, value=h); c.font = HEAD_FONT; c.fill = HEAD_FILL; c.alignment = WRAP; c.border = BORDER
for w, j in zip([44, 14, 18, 40, 60, 52, 44], range(1, 8)): ws.column_dimensions[get_column_letter(j)].width = w
row = 2
for g, n in cats.most_common():
    ex = next(r for r in S if cat(r) == g)
    vals = [g, n / N, round(TOTAL * n / N), RULES[g][0], RULES[g][1], ex['dns'] or ('(none) ' + ex['ip']), ex['os']]
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=row, column=j, value=v); c.font = TEXT; c.alignment = WRAP; c.border = BORDER
        if j == 2: c.number_format = '0.0%'
        if j == 3: c.number_format = '#,##0'
    row += 1
ws.cell(row=row + 1, column=1, value='Source: %d unmatched Discovered Items of the Qualys source; groups measured on a spread sample of %d items (twelve windows of 400 across the table), percentages applied to the full count.' % (TOTAL, N)).font = TEXT
ws.freeze_panes = 'A2'
ws2 = wb.create_sheet('Evidence carried')
rows2 = [('Hosts with a DNS name', sum(1 for r in S if r['dns'])), ('Hosts with a NetBIOS name', sum(1 for r in S if r['nb'])), ('Hosts with a Qualys host id', sum(1 for r in S if r['qg'])),
         ('Hosts with a serial number', sum(1 for r in S if r['sn'])), ('Hosts with a cloud resource id', sum(1 for r in S if r['cloud'])), ('Hosts scanned by IP (unauthenticated)', sum(1 for r in S if r['t'] == 'IP')),
         ('Hosts scanned by the Cloud Agent', sum(1 for r in S if r['t'] in ('AGENT', 'Cloud Agent'))), ('OS text empty', sum(1 for r in S if r['os'] in ('', '-', 'None'))), ('OS text is a multi-guess fingerprint', sum(1 for r in S if r['os'].count('/') >= 2))]
for j, h in enumerate(['Evidence', 'Share of sample', 'Estimated hosts'], 1):
    c = ws2.cell(row=1, column=j, value=h); c.font = HEAD_FONT; c.fill = HEAD_FILL; c.border = BORDER
for w, j in zip([44, 16, 18], range(1, 4)): ws2.column_dimensions[get_column_letter(j)].width = w
for i, (lab, n) in enumerate(rows2, 2):
    for j, v in enumerate([lab, n / N, round(TOTAL * n / N)], 1):
        c = ws2.cell(row=i, column=j, value=v); c.font = TEXT; c.border = BORDER
        if j == 2: c.number_format = '0.0%'
        if j == 3: c.number_format = '#,##0'
ws3 = wb.create_sheet('OS text (whole population)')
for j, h in enumerate(['OS text reported by Qualys', 'Unmatched hosts', 'Class the rules derive from it'], 1):
    c = ws3.cell(row=1, column=j, value=h); c.font = HEAD_FONT; c.fill = HEAD_FILL; c.border = BORDER
for w, j in zip([90, 16, 30], range(1, 4)): ws3.column_dimensions[get_column_letter(j)].width = w
def class_for(os_):
    s = os_.lower()
    if not s: return '(no OS text)'
    if s.count('/') >= 2: return '(multi-guess, no class)'
    for k, v in [('esx', 'ESX Server'), ('windows', 'Windows Server or Computer'), ('aix', 'AIX Server'), ('solaris', 'Solaris Server'), ('sunos', 'Solaris Server'), ('hp-ux', 'HPUX Server'), ('netapp', 'Storage Server'), ('ontap', 'Storage Server'), ('printer', 'Printer'), ('laserjet', 'Printer'), ('jetdirect', 'Printer')]:
        if k in s: return v
    if any(k in s for k in ('red hat', 'linux', 'centos', 'ubuntu', 'suse', 'debian', 'fedora', 'euleros', 'oracle enterprise', 'amazon')): return 'Linux Server'
    if any(k in s for k in ('nx-os', 'catos', 'cisco')): return 'Network Gear'
    return '(no class)'
for i, (os_, n) in enumerate(P['by_os'], 2):
    for j, v in enumerate([os_ or '(empty)', n, class_for(os_)], 1):
        c = ws3.cell(row=i, column=j, value=v); c.font = TEXT; c.border = BORDER
        if j == 2: c.number_format = '#,##0'
ws3.freeze_panes = 'A2'
out = os.path.join(HERE, 'Qualys Unmatched Hosts - Analysis.xlsx'); wb.save(out)
print('workbook written:', out); [print('  %-46s %5.1f%%  ~%7d' % (g, 100.0 * n / N, TOTAL * n / N)) for g, n in cats.most_common()]
